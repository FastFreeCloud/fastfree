from __future__ import annotations

from datetime import datetime

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

from .models import PatientData, TRACKING_COLUMNS, ARABIC_HEADERS, DEPARTMENTS

HEADER_FILL = PatternFill(start_color="003366", end_color="003366", fill_type="solid")
HEADER_FONT = Font(name="Arial", bold=True, color="FFFFFF", size=11)
SUBHEADER_FILL = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
SUBHEADER_FONT = Font(name="Arial", bold=True, color="FFFFFF", size=10)
DATA_FONT = Font(name="Arial", size=10)
THIN_BORDER = Border(
    left=Side(style="thin", color="B0B0B0"),
    right=Side(style="thin", color="B0B0B0"),
    top=Side(style="thin", color="B0B0B0"),
    bottom=Side(style="thin", color="B0B0B0"),
)
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
TRUE_FILL = PatternFill(start_color="D5F5E3", end_color="D5F5E3", fill_type="solid")
TRUE_FONT = Font(name="Arial", size=10, bold=True, color="27AE60")
DEPT_FILL = PatternFill(start_color="D6EAF8", end_color="D6EAF8", fill_type="solid")
DEPT_FONT = Font(name="Arial", bold=True, size=11, color="1A5276")

COLUMN_WIDTHS = {
    "A": 14, "B": 22, "C": 14, "D": 14, "E": 14, "F": 14,
    "G": 12, "H": 12, "I": 8, "J": 14, "K": 14, "L": 14,
    "M": 14, "N": 14, "O": 14, "P": 22,
}


def _flow_up_sheet(ws, records: list[PatientData], config: dict, departments=None, period="صباحي") -> None:
    for col_letter, width in COLUMN_WIDTHS.items():
        ws.column_dimensions[col_letter].width = width

    now = datetime.now()

    row_num = 1
    ws.merge_cells(start_row=row_num, start_column=1, end_row=row_num, end_column=16)
    c = ws.cell(row=row_num, column=1)
    c.value = "مستشفى منفلوط - شاشة المتابعة"
    c.font = Font(name="Arial", bold=True, size=14, color="003366")
    c.alignment = Alignment(horizontal="center", vertical="center")
    row_num += 1

    ws.merge_cells(start_row=row_num, start_column=1, end_row=row_num, end_column=16)
    c = ws.cell(row=row_num, column=1)
    c.value = f"{period} - {now.strftime('%Y-%m-%d')}"
    c.font = Font(name="Arial", bold=True, size=12, color="4472C4")
    c.alignment = Alignment(horizontal="center", vertical="center")
    row_num += 1

    for col_idx, value in enumerate(["BED", "NAME"] + ARABIC_HEADERS + ["المسؤول"], 1):
        c = ws.cell(row=row_num, column=col_idx)
        c.value = value
        c.fill = HEADER_FILL
        c.font = HEADER_FONT
        c.alignment = CENTER
        c.border = THIN_BORDER
    row_num += 1

    for col_idx, value in enumerate(["", ""] + TRACKING_COLUMNS + [""], 1):
        c = ws.cell(row=row_num, column=col_idx)
        c.value = value
        c.fill = SUBHEADER_FILL
        c.font = SUBHEADER_FONT
        c.alignment = CENTER
        c.border = THIN_BORDER
    row_num += 1

    ws.freeze_panes = ws.cell(row=row_num, column=1)

    grouped: dict[str, list[PatientData]] = {}
    for rec in records:
        dept = rec.bed.split("_")[0]
        grouped.setdefault(dept, []).append(rec)

    responsible_map = config.get("responsible", {})

    dept_cfg = config.get("departments")
    if departments is None:
        if dept_cfg is None:
            departments = [d.name for d in DEPARTMENTS]
        else:
            departments = [d.name for d in DEPARTMENTS if dept_cfg.get(d.name, True)]

    for dept_name in departments:
        dept_beds = grouped.get(dept_name, [])

        responsible = responsible_map.get(dept_name, "")

        for col_idx in range(1, 17):
            c = ws.cell(row=row_num, column=col_idx)
            c.fill = DEPT_FILL
            c.font = DEPT_FONT
            c.border = THIN_BORDER
            c.alignment = CENTER
        ws.cell(row=row_num, column=1).value = f"{dept_name} ({len(dept_beds)} beds)"
        ws.merge_cells(start_row=row_num, start_column=1, end_row=row_num, end_column=16)
        row_num += 1

        if not dept_beds:
            for col_idx in range(1, 17):
                c = ws.cell(row=row_num, column=col_idx)
                c.border = THIN_BORDER
                c.alignment = CENTER
            c = ws.cell(row=row_num, column=1)
            c.value = "لا توجد بيانات — تعذر تحميل القسم أو لا يوجد أسرّة"
            c.font = Font(name="Arial", size=10, italic=True, color="7F8C8D")
            ws.merge_cells(start_row=row_num, start_column=1, end_row=row_num, end_column=16)
            row_num += 1
            continue

        for rec in dept_beds:
            vals = ["TRUE" if v else "" for v in rec.to_values()]
            row_data = [rec.bed, rec.name] + vals + [responsible]
            for col_idx, value in enumerate(row_data, 1):
                c = ws.cell(row=row_num, column=col_idx)
                c.value = value
                c.border = THIN_BORDER
                if col_idx <= 2:
                    c.font = DATA_FONT
                    c.alignment = CENTER
                elif col_idx == 16:
                    c.font = Font(name="Arial", size=10, bold=True, color="003366")
                    c.alignment = CENTER
                elif value == "TRUE":
                    c.fill = TRUE_FILL
                    c.font = TRUE_FONT
                    c.alignment = CENTER
                else:
                    c.font = DATA_FONT
                    c.alignment = CENTER
            row_num += 1

    ws.row_dimensions[1].height = 30
    ws.row_dimensions[2].height = 25
    ws.row_dimensions[3].height = 35
    ws.row_dimensions[4].height = 30
    for r in range(5, row_num):
        ws.row_dimensions[r].height = 22


def generate_excel(records, output_dir, timestamp, config, departments=None, period="صباحي"):
    wb = Workbook()
    ws = wb.active
    ws.title = f"Flow Up {period}"
    _flow_up_sheet(ws, records, config, departments, period)
    path = output_dir / "FlowUp.xlsx"
    wb.save(path)
    return path
